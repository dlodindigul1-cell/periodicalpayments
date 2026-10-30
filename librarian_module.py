"""
நூலகர் இதழ் வருகைப் பதிவு (Librarian magazine-receipt register)
================================================================
app.py-ல் `init_librarian(app, get_conn)` என்று ஒரே ஒரு வரியில் இணைக்கப்படும் தனி Blueprint.

பகுதிகள்
  1. நூலகர் (Login dropdown-ல் "Librarian", Password இல்லை)
       /librarian/select      -> வகை + நூலகப் பெயர் தேர்வு
       /librarian/dashboard   -> இந்த வாரம் + முந்தைய வாரங்களின் (அதிகபட்சம் 4) "இன்னும் வராதவை" / "பதிவு செய்யாதவை": Received / Not received
  2. Admin
       /admin/librarian       -> Quarter Excel ஏற்றுதல், நூலகர் பதிவு நிலை (பதிவு செய்தவை / இன்னும் வராதவை / பதிவே செய்யாதவை),
                                 ஒரு நூலகத்தைத் தேர்ந்து Admin-ஆகவே Received / Not received பதிவு செய்தல்

Tables (தானாக உருவாகும்)
  lib_libraries (code, name, category)         <- data/libraries_seed.csv (காலியாக இருந்தால் ஒருமுறை) அல்லது Admin பக்கம்
  lib_schedule  (quarter, code, magazine, periodicity, expected_date)  <- Admin பக்கத்தில் Q3 / Q4 Excel ஏற்றுதல்
  lib_receipts  (code, magazine, expected_date, status, marked_at, marked_by)  <- நூலகர் / Admin பதிவு
  * lib_receipts தனி table என்பதால் Quarter Excel-ஐ மீண்டும் ஏற்றினாலும் நூலகர் பதிவுகள் அழியாது.
"""

import csv
import io
import os
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from functools import wraps

import psycopg2.extras
from flask import Blueprint, Response, jsonify, redirect, render_template_string, request, session

lib_bp = Blueprint("librarian", __name__)
_get_conn = None
_ready = False

IST = timezone(timedelta(hours=5, minutes=30))
SEED_CSV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "libraries_seed.csv")
STATUSES = ("Received", "Not received")
CARRY_WEEKS = 4   # நூலகர் திரையில் முந்தைய வாரங்களின் இதழ்கள் அதிகபட்சம் எத்தனை வாரம் காட்டப்படும்
CATEGORY_ORDER = ["மாவட்ட மைய நூலகம்", "முழுநேர கிளை நூலகம்", "கிளை நூலகம்", "ஊர்ப்புற நூலகம்"]
TA_DAYS = ["திங்கள்", "செவ்வாய்", "புதன்", "வியாழன்", "வெள்ளி", "சனி", "ஞாயிறு"]


def init_librarian(app, get_conn):
    global _get_conn
    _get_conn = get_conn
    app.register_blueprint(lib_bp)


# --------------------------------------------------------------------------- #
# பொது உதவிகள்
# --------------------------------------------------------------------------- #
def today_ist():
    return datetime.now(IST).date()


def week_bounds(d):
    """திங்கள் முதல் ஞாயிறு வரை."""
    start = d - timedelta(days=d.weekday())
    return start, start + timedelta(days=6)


def quarter_label(d):
    """நிதியாண்டு Quarter: Apr-Jun=Q1, Jul-Sep=Q2, Oct-Dec=Q3, Jan-Mar=Q4 (எ.கா. 2026-2027-Q3)."""
    fy = d.year if d.month >= 4 else d.year - 1
    q = {4: 1, 5: 1, 6: 1, 7: 2, 8: 2, 9: 2, 10: 3, 11: 3, 12: 3, 1: 4, 2: 4, 3: 4}[d.month]
    return f"{fy}-{fy + 1}-Q{q}"


def fmt_d(d):
    return d.strftime("%d/%m/%Y") if d else ""


def parse_iso(s):
    try:
        return datetime.strptime((s or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def ensure_schema():
    global _ready
    if _ready:
        return
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT pg_advisory_xact_lock(7712001)")   # பல worker ஒரே நேரத்தில் உருவாக்குவதைத் தடுக்க
            cur.execute(
                "CREATE TABLE IF NOT EXISTS lib_libraries ("
                " code TEXT PRIMARY KEY, name TEXT NOT NULL, category TEXT NOT NULL DEFAULT '',"
                " updated_at TIMESTAMPTZ NOT NULL DEFAULT now())")
            cur.execute(
                "CREATE TABLE IF NOT EXISTS lib_schedule ("
                " id SERIAL PRIMARY KEY, quarter TEXT NOT NULL, code TEXT NOT NULL, magazine TEXT NOT NULL,"
                " periodicity TEXT, expected_date DATE NOT NULL,"
                " UNIQUE (code, magazine, expected_date))")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_lib_sched_date ON lib_schedule (expected_date)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_lib_sched_quarter ON lib_schedule (quarter)")
            cur.execute(
                "CREATE TABLE IF NOT EXISTS lib_receipts ("
                " code TEXT NOT NULL, magazine TEXT NOT NULL, expected_date DATE NOT NULL,"
                " status TEXT NOT NULL CHECK (status IN ('Received','Not received')),"
                " marked_at TIMESTAMPTZ NOT NULL DEFAULT now(),"
                " PRIMARY KEY (code, magazine, expected_date))")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_lib_receipts_date ON lib_receipts (expected_date)")
            # யார் பதிவு செய்தார்: librarian / admin (பழைய பதிவுகள் = librarian)
            cur.execute("ALTER TABLE lib_receipts ADD COLUMN IF NOT EXISTS marked_by TEXT NOT NULL DEFAULT 'librarian'")
            cur.execute("SELECT COUNT(*) AS n FROM lib_libraries")
            if cur.fetchone()["n"] == 0 and os.path.exists(SEED_CSV):
                with open(SEED_CSV, newline="", encoding="utf-8-sig") as f:
                    rows = [(r["code"].strip(), r["name"].strip(), r["category"].strip())
                            for r in csv.DictReader(f) if (r.get("code") or "").strip()]
                psycopg2.extras.execute_values(
                    cur, "INSERT INTO lib_libraries (code, name, category) VALUES %s ON CONFLICT (code) DO NOTHING", rows)
        conn.commit()
    finally:
        conn.close()
    _ready = True


def _user():
    return session.get("user") or {}


def librarian_page(f):
    @wraps(f)
    def w(*a, **k):
        if _user().get("role") != "librarian":
            return redirect("/")
        ensure_schema()
        return f(*a, **k)
    return w


def librarian_api(need_library=True):
    def deco(f):
        @wraps(f)
        def w(*a, **k):
            if _user().get("role") != "librarian":
                return jsonify({"success": False, "message": "நூலகர் நுழைவு தேவை."}), 403
            ensure_schema()
            if need_library and not session.get("lib_code"):
                return jsonify({"success": False, "needLibrary": True, "message": "முதலில் நூலகத்தைத் தேர்ந்தெடுக்கவும்."}), 400
            return f(*a, **k)
        return w
    return deco


def admin_page(f):
    @wraps(f)
    def w(*a, **k):
        if _user().get("role") != "admin":
            return redirect("/")
        ensure_schema()
        return f(*a, **k)
    return w


def admin_api(f):
    @wraps(f)
    def w(*a, **k):
        if _user().get("role") != "admin":
            return jsonify({"success": False, "message": "இந்த வசதி Admin-க்கு மட்டுமே."}), 403
        ensure_schema()
        return f(*a, **k)
    return w


def _category_sort_key(c):
    return (CATEGORY_ORDER.index(c) if c in CATEGORY_ORDER else len(CATEGORY_ORDER), c)


# =========================================================================== #
# 1. நூலகர் பக்கங்கள்
# =========================================================================== #
@lib_bp.route("/librarian/select")
@librarian_page
def librarian_select_page():
    return render_template_string(SELECT_HTML)


@lib_bp.route("/librarian/change")
@librarian_page
def librarian_change():
    session.pop("lib_code", None)
    session.pop("lib_name", None)
    return redirect("/librarian/select")


def _libraries_json():
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT code, name, category FROM lib_libraries")
            rows = cur.fetchall()
    finally:
        conn.close()
    groups = {}
    for r in rows:
        groups.setdefault(r["category"] or "மற்றவை", []).append({"code": r["code"], "name": r["name"]})
    out = [{"category": c, "libraries": sorted(v, key=lambda x: x["name"])}
           for c, v in sorted(groups.items(), key=lambda kv: _category_sort_key(kv[0]))]
    return jsonify({"success": True, "categories": out})


@lib_bp.route("/api/librarian/libraries")
@librarian_api(need_library=False)
def api_librarian_libraries():
    return _libraries_json()


@lib_bp.route("/api/librarian/select", methods=["POST"])
@librarian_api(need_library=False)
def api_librarian_select():
    code = ((request.get_json(force=True, silent=True) or {}).get("code") or "").strip()
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT code, name FROM lib_libraries WHERE code=%s", (code,))
            row = cur.fetchone()
    finally:
        conn.close()
    if not row:
        return jsonify({"success": False, "message": "நூலகம் கண்டறியப்படவில்லை."}), 404
    session["lib_code"], session["lib_name"] = row["code"], row["name"]
    return jsonify({"success": True})


@lib_bp.route("/librarian/dashboard")
@librarian_page
def librarian_dashboard_page():
    if not session.get("lib_code"):
        return redirect("/librarian/select")
    return render_template_string(DASH_HTML)


def _row_json(r, ws):
    d = r["expected_date"]
    return {"magazine": r["magazine"], "periodicity": r["periodicity"] or "",
            "date": d.isoformat(), "dateLabel": fmt_d(d), "day": TA_DAYS[d.weekday()],
            "status": r["status"], "weeksAgo": (ws - week_bounds(d)[0]).days // 7}


def _library_week_payload(code, weeks_back):
    """ஒரு நூலகத்துக்கு: இந்த வாரம் + முந்தைய வாரங்களில் (weeks_back வாரம் வரை; None = வரம்பின்றி)
    'இன்னும் வராதவை' (Not received) மற்றும் 'பதிவு செய்யாதவை'. நூலகர் திரையும் Admin பதிவுத் திரையும் இதையே பயன்படுத்தும்."""
    today = today_ist()
    ws, we = week_bounds(today)
    join = ("LEFT JOIN lib_receipts r ON r.code=s.code AND r.magazine=s.magazine "
            "AND r.expected_date=s.expected_date ")
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT code, name, category FROM lib_libraries WHERE code=%s", (code,))
            lib = cur.fetchone()
            if not lib:
                return None
            cur.execute(
                "SELECT s.magazine, s.periodicity, s.expected_date, r.status FROM lib_schedule s " + join +
                "WHERE s.code=%s AND s.expected_date BETWEEN %s AND %s ORDER BY s.expected_date, s.magazine",
                (code, ws, we))
            this_rows = cur.fetchall()
            sql = ("SELECT s.magazine, s.periodicity, s.expected_date, r.status FROM lib_schedule s " + join +
                   "WHERE s.code=%s AND s.expected_date < %s")
            params = [code, ws]
            if weeks_back:
                sql += " AND s.expected_date >= %s"
                params.append(ws - timedelta(weeks=weeks_back))
            sql += " AND (r.code IS NULL OR r.status='Not received') ORDER BY s.expected_date, s.magazine"
            cur.execute(sql, params)
            earlier = cur.fetchall()
    finally:
        conn.close()
    return {
        "success": True, "library": dict(lib),
        "weekStart": fmt_d(ws), "weekEnd": fmt_d(we), "today": today.isoformat(), "carryWeeks": weeks_back or 0,
        "rows": [_row_json(r, ws) for r in this_rows],
        "earlierNotReceived": [_row_json(r, ws) for r in earlier if r["status"] == "Not received"],
        "earlierUnregistered": [_row_json(r, ws) for r in earlier if not r["status"]],
    }


def _save_entries(code, entries, weeks_back, by):
    """Received / Not received சேமிப்பு. எதிர்காலத் தேதி (இந்த வாரத்துக்குப் பிந்தையவை) எப்போதும் தடை.
    weeks_back: இந்த வாரத்துக்கு முந்தைய எத்தனை வாரம் வரை அனுமதி (None = வரம்பின்றி — Admin)."""
    ws, we = week_bounds(today_ist())
    lo = ws - timedelta(weeks=weeks_back) if weeks_back else None
    clean = {}
    for e in entries:
        d = parse_iso(e.get("date"))
        mag, st = (e.get("magazine") or "").strip(), e.get("status")
        if not mag or st not in STATUSES or not d or d > we or (lo and d < lo):
            return jsonify({"success": False,
                            "message": "தவறான பதிவு (இந்த வாரம் மற்றும் அதற்கு முந்தையவை மட்டுமே; எதிர்காலத் தேதிகள் முடியாது)."}), 400
        clean[(mag, d)] = (code, mag, d, st, by)
    rows = list(clean.values())
    if not rows:
        return jsonify({"success": False, "message": "எந்த இதழும் தேர்ந்தெடுக்கப்படவில்லை."}), 400
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            # அட்டவணையில் உள்ள இதழ்களை மட்டுமே ஏற்கும்
            cur.execute("SELECT magazine, expected_date FROM lib_schedule WHERE code=%s AND expected_date BETWEEN %s AND %s",
                        (code, min(r[2] for r in rows), max(r[2] for r in rows)))
            valid = {(r["magazine"], r["expected_date"]) for r in cur.fetchall()}
            rows = [r for r in rows if (r[1], r[2]) in valid]
            if not rows:
                return jsonify({"success": False, "message": "பதிவு செய்ய இதழ்கள் எதுவும் பொருந்தவில்லை."}), 400
            psycopg2.extras.execute_values(
                cur,
                "INSERT INTO lib_receipts (code, magazine, expected_date, status, marked_by) VALUES %s "
                "ON CONFLICT (code, magazine, expected_date) DO UPDATE "
                "SET status=EXCLUDED.status, marked_at=now(), marked_by=EXCLUDED.marked_by "
                "WHERE lib_receipts.status IS DISTINCT FROM EXCLUDED.status",
                rows)
        conn.commit()
    except Exception as e:  # noqa: BLE001
        conn.rollback()
        return jsonify({"success": False, "message": str(e)}), 500
    finally:
        conn.close()
    who = "Admin மூலம் " if by == "admin" else ""
    return jsonify({"success": True, "saved": len(rows), "message": f"✅ {len(rows)} இதழ்களின் பதிவு {who}சேமிக்கப்பட்டது."})


@lib_bp.route("/api/librarian/week")
@librarian_api()
def api_librarian_week():
    data = _library_week_payload(session["lib_code"], CARRY_WEEKS)
    if data is None:
        return jsonify({"success": False, "needLibrary": True, "message": "நூலகம் கண்டறியப்படவில்லை."}), 400
    return jsonify(data)


@lib_bp.route("/api/librarian/save", methods=["POST"])
@librarian_api()
def api_librarian_save():
    entries = (request.get_json(force=True, silent=True) or {}).get("entries") or []
    return _save_entries(session["lib_code"], entries, CARRY_WEEKS, "librarian")


# ---- ஆணையிடப்பட்ட இதழ்கள் பட்டியல் (PDF / Excel) ------------------------------------------------ #
HEAD_LINES = ["தமிழ்நாடு அரசு பொது நூலகத்துறை", "மாவட்ட நூலக ஆணைக்குழு, திண்டுக்கல்"]
REG_TITLE = "பருவ இதழ் வருகைப் பதிவேடு"
ORDERED_TITLE = "உங்கள் நூலகத்திற்கு ஆணையிடப்பட்ட இதழ்கள்"
_FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "fonts")


def _is_tamil(text):
    """இதழ் பெயரில் தமிழ் எழுத்து (U+0B80–U+0BFF) இருந்தால் தமிழ் இதழ்."""
    return any("\u0b80" <= ch <= "\u0bff" for ch in (text or ""))


def _ordered_magazines(code):
    """(நூலகப் பெயர், ஆங்கில இதழ்கள், தமிழ் இதழ்கள்) — ஒவ்வொன்றும் [(இதழ், காலமுறை)].
    நடப்பு Quarter-ல் அட்டவணை இருந்தால் அது; இல்லையெனில் அந்த நூலகத்தின் எல்லா அட்டவணையும்."""
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT name FROM lib_libraries WHERE code=%s", (code,))
            lib = cur.fetchone()
            if not lib:
                return None
            cur.execute("SELECT magazine, MAX(periodicity) AS periodicity FROM lib_schedule "
                        "WHERE code=%s AND quarter=%s GROUP BY magazine", (code, quarter_label(today_ist())))
            rows = cur.fetchall()
            if not rows:
                cur.execute("SELECT magazine, MAX(periodicity) AS periodicity FROM lib_schedule "
                            "WHERE code=%s GROUP BY magazine", (code,))
                rows = cur.fetchall()
    finally:
        conn.close()
    items = sorted(((r["magazine"], r["periodicity"] or "") for r in rows), key=lambda x: x[0].casefold())
    return (lib["name"], [i for i in items if not _is_tamil(i[0])], [i for i in items if _is_tamil(i[0])])


def _ordered_pdf(lib_name, eng, tam):
    from fpdf import FPDF
    from fpdf.fonts import FontFace
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("Tamil", "", os.path.join(_FONT_DIR, "NotoSansTamil-Regular.ttf"))
    pdf.add_font("Tamil", "B", os.path.join(_FONT_DIR, "NotoSansTamil-Bold.ttf"))
    pdf.set_text_shaping(True)
    pdf.add_page()
    for txt, size in [(HEAD_LINES[0], 15), (HEAD_LINES[1], 13), (lib_name, 13), (REG_TITLE, 13)]:
        pdf.set_font("Tamil", "B", size)
        pdf.multi_cell(0, 7.5, txt, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Tamil", "B", 12)
    pdf.multi_cell(0, 7, ORDERED_TITLE, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Tamil", "", 11)
    head = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(30, 77, 140))
    sec = FontFace(emphasis="BOLD", fill_color=(220, 230, 247))
    with pdf.table(col_widths=(22, 108, 50), text_align=("CENTER", "LEFT", "CENTER"), headings_style=head,
                   line_height=7, padding=1.5, borders_layout="ALL") as t:
        h = t.row()
        for c in ("வ.எண்.", "இதழ் பெயர்", "காலமுறை"):
            h.cell(c)
        for title, items in (("ஆங்கிலம்", eng), ("தமிழ்", tam)):
            r = t.row()
            r.cell(title, colspan=3, align="L", style=sec)
            if not items:
                t.row().cell("— இல்லை —", colspan=3, align="C")
            for n, (mag, per) in enumerate(items, start=1):
                r = t.row()
                r.cell(str(n))
                r.cell(mag)
                r.cell(per)
    return bytes(pdf.output())


def _ordered_xlsx(lib_name, eng, tam):
    import xlsxwriter
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("Ordered magazines")
    big = wb.add_format({"bold": True, "align": "center", "valign": "vcenter", "font_size": 14})
    mid = wb.add_format({"bold": True, "align": "center", "valign": "vcenter", "font_size": 12})
    head = wb.add_format({"bold": True, "bg_color": "#1e4d8c", "font_color": "#ffffff", "border": 1, "align": "center"})
    sec = wb.add_format({"bold": True, "bg_color": "#dce6f7", "border": 1})
    cell = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
    mid_c = wb.add_format({"border": 1, "align": "center", "valign": "top"})
    r = 0
    for txt, f in [(HEAD_LINES[0], big), (HEAD_LINES[1], mid), (lib_name, mid), (REG_TITLE, mid), (ORDERED_TITLE, mid)]:
        ws.merge_range(r, 0, r, 2, txt, f)
        r += 1
    r += 1
    for c, h in enumerate(("வ.எண்.", "இதழ் பெயர்", "காலமுறை")):
        ws.write(r, c, h, head)
    r += 1
    for title, items in (("ஆங்கிலம்", eng), ("தமிழ்", tam)):
        ws.merge_range(r, 0, r, 2, title, sec)
        r += 1
        for n, (mag, per) in enumerate(items, start=1):
            ws.write_number(r, 0, n, mid_c)
            ws.write(r, 1, mag, cell)
            ws.write(r, 2, per, mid_c)
            r += 1
    ws.set_column(0, 0, 10)
    ws.set_column(1, 1, 60)
    ws.set_column(2, 2, 22)
    wb.close()
    return buf.getvalue()


@lib_bp.route("/api/librarian/ordered-export")
@librarian_api()
def api_librarian_ordered_export():
    code = session["lib_code"]
    data = _ordered_magazines(code)
    if data is None:
        return jsonify({"success": False, "message": "நூலகம் கண்டறியப்படவில்லை."}), 400
    lib_name, eng, tam = data
    fmt = (request.args.get("fmt") or "pdf").lower()
    stamp = today_ist().isoformat()
    if fmt in ("xlsx", "excel"):
        body, mime, ext = _ordered_xlsx(lib_name, eng, tam), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        body, mime, ext = _ordered_pdf(lib_name, eng, tam), "application/pdf", "pdf"
    return Response(body, mimetype=mime,
                    headers={"Content-Disposition": f"attachment; filename=ordered_magazines_{code}_{stamp}.{ext}"})


# ---- இந்த மாதம் / முந்தைய மாதம்: வந்தவை - வராதவை ---------------------------------------------------- #
TA_MONTHS = ["ஜனவரி", "பிப்ரவரி", "மார்ச்", "ஏப்ரல்", "மே", "ஜூன்",
             "ஜூலை", "ஆகஸ்ட்", "செப்டம்பர்", "அக்டோபர்", "நவம்பர்", "டிசம்பர்"]


def month_bounds(d, offset=0):
    """d-ன் மாதம் (offset=1 -> முந்தைய மாதம்) முதல் நாள், கடைசி நாள்."""
    first = d.replace(day=1)
    for _ in range(offset):
        first = (first - timedelta(days=1)).replace(day=1)
    nxt = (first + timedelta(days=32)).replace(day=1)
    return first, nxt - timedelta(days=1)


@lib_bp.route("/api/librarian/month")
@librarian_api()
def api_librarian_month():
    offset = 1 if request.args.get("offset") == "1" else 0
    first, last = month_bounds(today_ist(), offset)
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT s.magazine, s.expected_date, r.status FROM lib_schedule s "
                "JOIN lib_receipts r ON r.code=s.code AND r.magazine=s.magazine AND r.expected_date=s.expected_date "
                "WHERE s.code=%s AND s.expected_date BETWEEN %s AND %s "
                "ORDER BY s.expected_date DESC, s.magazine", (session["lib_code"], first, last))
            rows = cur.fetchall()
    finally:
        conn.close()
    pick = lambda st: [{"magazine": r["magazine"], "dateLabel": fmt_d(r["expected_date"])} for r in rows if r["status"] == st]
    return jsonify({"success": True, "offset": offset, "label": f"{TA_MONTHS[first.month - 1]} {first.year}",
                    "received": pick("Received"), "notReceived": pick("Not received")})


# =========================================================================== #
# 2. Excel / CSV வாசிப்பு (Quarter ஏற்றுதல்)
# =========================================================================== #
def _iter_sheet_rows(fs):
    """.xlsx / .xls / .csv கோப்பின் வரிசைகளை tuple-ஆகத் தரும் (நெடுவரிசை எண்ணின்படி; தலைப்புப் பெயர் முக்கியமில்லை)."""
    name = (fs.filename or "").lower()
    data = fs.read()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            yield row
    elif name.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        sh = wb.sheet_by_index(0)
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
                else:
                    row.append(cell.value)
            yield tuple(row)
    elif name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        for row in csv.reader(io.StringIO(text)):
            yield tuple(row)
    else:
        raise ValueError("கோப்பு .xlsx, .xls அல்லது .csv ஆக இருக்க வேண்டும்.")


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        from openpyxl.utils.datetime import from_excel
        return from_excel(v).date()
    s = str(v or "").strip()
    if not s:
        return None
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%b-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# =========================================================================== #
# 3. Admin பக்கம் + API-கள்
# =========================================================================== #
@lib_bp.route("/admin/librarian")
@admin_page
def admin_librarian_page():
    return render_template_string(ADMIN_HTML)


@lib_bp.route("/api/admin/librarian/overview")
@admin_api
def api_admin_overview():
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS n FROM lib_libraries")
            libs = cur.fetchone()["n"]
            cur.execute("SELECT DISTINCT category FROM lib_libraries ORDER BY category")
            cats = sorted([r["category"] for r in cur.fetchall() if r["category"]], key=_category_sort_key)
            cur.execute(
                "SELECT s.quarter, COUNT(*) AS rows, MIN(s.expected_date) AS dfrom, MAX(s.expected_date) AS dto, "
                "       COUNT(DISTINCT s.code) AS libs, COUNT(DISTINCT s.magazine) AS mags, "
                "       COUNT(r.code) AS recorded "
                "FROM lib_schedule s LEFT JOIN lib_receipts r "
                "  ON r.code=s.code AND r.magazine=s.magazine AND r.expected_date=s.expected_date "
                "GROUP BY s.quarter ORDER BY MIN(s.expected_date)")
            quarters = [{"quarter": r["quarter"], "rows": r["rows"], "from": fmt_d(r["dfrom"]), "to": fmt_d(r["dto"]),
                         "libs": r["libs"], "mags": r["mags"], "recorded": r["recorded"]} for r in cur.fetchall()]
    finally:
        conn.close()
    ws, we = week_bounds(today_ist())
    return jsonify({"success": True, "libraries": libs, "categories": cats, "quarters": quarters,
                    "weekStart": ws.isoformat(), "weekEnd": we.isoformat(), "today": today_ist().isoformat()})


@lib_bp.route("/api/admin/librarian/upload-quarter", methods=["POST"])
@admin_api
def api_admin_upload_quarter():
    """Quarter Excel (A: Library code, B: Magazine, C: Periodicity, D: Expected Date).
    Quarter தேதியிலிருந்து தானாகக் கணிக்கப்படும். mode=replace (இயல்பு): கோப்பிலுள்ள Quarter-களின் பழைய அட்டவணை மாற்றப்படும்;
    நூலகர் பதிவுகள் (lib_receipts) பாதிக்கப்படாது. mode=add: புதியவை மட்டும் சேர்க்கப்படும்/புதுப்பிக்கப்படும்."""
    fs = request.files.get("file")
    if not fs or not fs.filename:
        return jsonify({"success": False, "message": "கோப்பைத் தேர்ந்தெடுக்கவும்."}), 400
    mode = request.form.get("mode", "replace")
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT code FROM lib_libraries")
            known = {r["code"] for r in cur.fetchall()}
        if not known:
            return jsonify({"success": False, "message": "நூலகப் பட்டியல் (code) இன்னும் இல்லை — முதலில் அதை ஏற்றவும்."}), 400
        rows, seen = [], set()
        unknown, bad, dup, bad_samples = Counter(), 0, 0, []
        try:
            for i, row in enumerate(_iter_sheet_rows(fs)):
                if not row or len(row) < 4:
                    continue
                code, mag, per, d = _txt(row[0]), _txt(row[1]), _txt(row[2]), _to_date(row[3])
                if i == 0 and d is None:
                    continue                      # தலைப்பு வரிசை
                if not code and not mag and d is None:
                    continue                      # வெற்று வரிசை
                if not code or not mag or d is None:
                    bad += 1
                    if len(bad_samples) < 5:
                        bad_samples.append(i + 1)
                    continue
                if code not in known:
                    unknown[code] += 1
                    continue
                key = (code, mag, d)
                if key in seen:
                    dup += 1
                    continue
                seen.add(key)
                rows.append((quarter_label(d), code, mag, per, d))
        except ValueError as e:
            return jsonify({"success": False, "message": str(e)}), 400
        except ImportError:
            return jsonify({"success": False, "message": ".xls படிக்க xlrd தேவை (requirements.txt-ல் xlrd உள்ளது; மீண்டும் Deploy செய்யவும்)."}), 500
        if not rows:
            return jsonify({"success": False, "message": "சரியான வரிசைகள் எதுவும் இல்லை. நெடுவரிசை வரிசை: Library code, Magazine, Periodicity, Expected Date."}), 400
        quarters = sorted({r[0] for r in rows})
        with conn.cursor() as cur:
            if mode == "replace":
                cur.execute("DELETE FROM lib_schedule WHERE quarter = ANY(%s)", (quarters,))
            psycopg2.extras.execute_values(
                cur,
                "INSERT INTO lib_schedule (quarter, code, magazine, periodicity, expected_date) VALUES %s "
                "ON CONFLICT (code, magazine, expected_date) DO UPDATE "
                "SET periodicity=EXCLUDED.periodicity, quarter=EXCLUDED.quarter",
                rows, page_size=2000)
        conn.commit()
    except Exception as e:  # noqa: BLE001
        conn.rollback()
        return jsonify({"success": False, "message": "பிழை: " + str(e)}), 500
    finally:
        conn.close()
    per_q = Counter(r[0] for r in rows)
    warn = []
    if unknown:
        warn.append(f"நூலகப் பட்டியலில் இல்லாத code உள்ள {sum(unknown.values())} வரிசைகள் தவிர்க்கப்பட்டன ({', '.join(list(unknown)[:5])}…)")
    if bad:
        warn.append(f"தேதி / இதழ் / code குறைபாடுள்ள {bad} வரிசைகள் தவிர்க்கப்பட்டன (வரிசை எண்: {', '.join(map(str, bad_samples))}…)")
    if dup:
        warn.append(f"இரட்டை வரிசைகள் {dup} புறக்கணிக்கப்பட்டன")
    return jsonify({
        "success": True,
        "message": f"✅ {len(rows)} வரிசைகள் ஏற்றப்பட்டன — " + ", ".join(f"{q}: {n}" for q, n in sorted(per_q.items())),
        "warnings": warn,
        "libraries": len({r[1] for r in rows}), "magazines": len({r[2] for r in rows}),
    })


@lib_bp.route("/api/admin/librarian/delete-quarter", methods=["POST"])
@admin_api
def api_admin_delete_quarter():
    """தவறான Quarter ஏற்றத்தை நீக்க: அந்த Quarter அட்டவணை + அந்த தேதிகளுக்கான நூலகர் பதிவுகளும் நீக்கப்படும்."""
    q = ((request.get_json(force=True, silent=True) or {}).get("quarter") or "").strip()
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT MIN(expected_date) AS a, MAX(expected_date) AS b FROM lib_schedule WHERE quarter=%s", (q,))
            r = cur.fetchone()
            if not r or not r["a"]:
                return jsonify({"success": False, "message": "அந்த Quarter இல்லை."}), 404
            cur.execute("DELETE FROM lib_receipts WHERE expected_date BETWEEN %s AND %s", (r["a"], r["b"]))
            nr = cur.rowcount
            cur.execute("DELETE FROM lib_schedule WHERE quarter=%s", (q,))
            ns = cur.rowcount
        conn.commit()
    finally:
        conn.close()
    return jsonify({"success": True, "message": f"🗑️ {q}: அட்டவணை {ns} வரிசைகள், நூலகர் பதிவுகள் {nr} நீக்கப்பட்டன."})


@lib_bp.route("/api/admin/librarian/upload-libraries", methods=["POST"])
@admin_api
def api_admin_upload_libraries():
    """நூலகப் பட்டியல்: name, category, code (தலைப்புப் பெயர் இருந்தால் அதன்படி; இல்லையெனில் இதே வரிசை). code வாரியாக UPSERT."""
    fs = request.files.get("file")
    if not fs or not fs.filename:
        return jsonify({"success": False, "message": "கோப்பைத் தேர்ந்தெடுக்கவும்."}), 400
    try:
        data = list(_iter_sheet_rows(fs))
    except ValueError as e:
        return jsonify({"success": False, "message": str(e)}), 400
    except ImportError:
        return jsonify({"success": False, "message": ".xls படிக்க xlrd தேவை."}), 500
    idx = {"name": 0, "category": 1, "code": 2}
    if data:
        head = [_txt(c).lower() for c in data[0]]
        if "code" in head and "name" in head:
            idx = {"name": head.index("name"), "category": head.index("category") if "category" in head else -1,
                   "code": head.index("code")}
            data = data[1:]
        elif head and head[0] in ("name", "library", "library name"):
            data = data[1:]
    rows = {}
    for r in data:
        code = _txt(r[idx["code"]]) if len(r) > idx["code"] else ""
        name = _txt(r[idx["name"]]) if len(r) > idx["name"] else ""
        cat = _txt(r[idx["category"]]) if idx["category"] >= 0 and len(r) > idx["category"] else ""
        if code and name:
            rows[code] = (code, name, cat)
    if not rows:
        return jsonify({"success": False, "message": "சரியான வரிசைகள் இல்லை (name, category, code)."}), 400
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            psycopg2.extras.execute_values(
                cur,
                "INSERT INTO lib_libraries (code, name, category) VALUES %s ON CONFLICT (code) DO UPDATE "
                "SET name=EXCLUDED.name, category=EXCLUDED.category, updated_at=now()", list(rows.values()))
        conn.commit()
    finally:
        conn.close()
    return jsonify({"success": True, "message": f"✅ {len(rows)} நூலகங்கள் சேமிக்கப்பட்டன/புதுப்பிக்கப்பட்டன."})


# ---- Admin: பதிவு செய்தவை / இன்னும் வராதவை / பதிவே செய்யாதவை ------------------------------------- #
STATUS_TABS = ("received", "notreceived", "unregistered")


def _delay_days(marked_at, expected):
    """எதிர்பார்த்த தேதிக்குப் பிறகு Received எனக் குறித்த நாட்கள் (தாமதம்)."""
    if not marked_at or not expected:
        return 0
    return max((marked_at.astimezone(IST).date() - expected).days, 0)


def _filters():
    """status (received | notreceived | unregistered), category, from, to, code.
    'unregistered'-க்கு 'to' <= இன்று."""
    status = request.args.get("status", "received")
    if status not in STATUS_TABS:
        status = "received"
    cat = (request.args.get("category") or "").strip()
    dfrom, dto = parse_iso(request.args.get("from")), parse_iso(request.args.get("to"))
    code = (request.args.get("code") or "").strip()
    where, params = [], []
    if cat:
        where.append("l.category = %s")
        params.append(cat)
    if code:
        where.append("l.code = %s")
        params.append(code)
    if status == "received":
        where.append("r.status = 'Received'")
    elif status == "notreceived":
        where.append("r.status = 'Not received'")
    else:
        cap = today_ist()
        dto = min(dto, cap) if dto else cap
    dcol = "s.expected_date" if status == "unregistered" else "r.expected_date"
    if dfrom:
        where.append(dcol + " >= %s")
        params.append(dfrom)
    if dto:
        where.append(dcol + " <= %s")
        params.append(dto)
    return status, where, params


_TZ_DAYS = "((r.marked_at AT TIME ZONE 'Asia/Kolkata')::date - r.expected_date)"


@lib_bp.route("/api/admin/librarian/summary")
@admin_api
def api_admin_summary():
    status, where, params = _filters()
    extra = (" AND " + " AND ".join(where)) if where else ""
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS n FROM lib_libraries")
            total_libs = cur.fetchone()["n"]
            if status == "unregistered":
                cur.execute(
                    "SELECT l.code, l.name, l.category, COUNT(*) AS cnt, COUNT(DISTINCT s.magazine) AS mags, "
                    "       MIN(s.expected_date) AS first_d, "
                    "       (SELECT COUNT(*) FROM lib_receipts rr WHERE rr.code=l.code) AS recorded_total "
                    "FROM lib_schedule s JOIN lib_libraries l ON l.code=s.code "
                    "LEFT JOIN lib_receipts r ON r.code=s.code AND r.magazine=s.magazine AND r.expected_date=s.expected_date "
                    "WHERE r.code IS NULL" + extra + " GROUP BY l.code, l.name, l.category ORDER BY l.name", params)
                items = [{"code": r["code"], "name": r["name"], "category": r["category"], "count": r["cnt"],
                          "magazines": r["mags"], "firstDue": fmt_d(r["first_d"]), "recordedTotal": r["recorded_total"]}
                         for r in cur.fetchall()]
            else:
                cur.execute(
                    "SELECT l.code, l.name, l.category, COUNT(*) AS cnt, COUNT(DISTINCT r.magazine) AS mags, "
                    "       MIN(r.expected_date) AS first_d, MAX(r.marked_at) AS last_at, "
                    f"       SUM(CASE WHEN {_TZ_DAYS} > 0 THEN 1 ELSE 0 END) AS late, "
                    f"       MAX(GREATEST({_TZ_DAYS}, 0)) AS max_late, "
                    "       SUM(CASE WHEN r.marked_by = 'admin' THEN 1 ELSE 0 END) AS by_admin "
                    "FROM lib_receipts r JOIN lib_libraries l ON l.code=r.code "
                    "WHERE TRUE" + extra + " GROUP BY l.code, l.name, l.category ORDER BY l.name", params)
                items = [{"code": r["code"], "name": r["name"], "category": r["category"], "count": r["cnt"],
                          "magazines": r["mags"], "firstDue": fmt_d(r["first_d"]),
                          "late": int(r["late"] or 0), "maxLate": int(r["max_late"] or 0),
                          "byAdmin": int(r["by_admin"] or 0),
                          "lastAt": r["last_at"].astimezone(IST).strftime("%d/%m/%Y %H:%M") if r["last_at"] else ""}
                         for r in cur.fetchall()]
            cur.execute("SELECT COUNT(DISTINCT code) AS n FROM lib_receipts")
            ever = cur.fetchone()["n"]
    finally:
        conn.close()
    return jsonify({"success": True, "items": items, "totalLibraries": total_libs, "librariesEverRecorded": ever,
                    "totalRows": sum(i["count"] for i in items), "lateTotal": sum(i.get("late", 0) for i in items)})


def _detail_rows(limit=None):
    status, where, params = _filters()
    extra = (" AND " + " AND ".join(where)) if where else ""
    lim = f" LIMIT {int(limit)}" if limit else ""
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            if status == "unregistered":
                cur.execute(
                    "SELECT l.code, l.name, l.category, s.magazine, s.periodicity, s.expected_date, NULL::text AS status, "
                    "       NULL::timestamptz AS marked_at, NULL::text AS marked_by "
                    "FROM lib_schedule s JOIN lib_libraries l ON l.code=s.code "
                    "LEFT JOIN lib_receipts r ON r.code=s.code AND r.magazine=s.magazine AND r.expected_date=s.expected_date "
                    "WHERE r.code IS NULL" + extra + " ORDER BY l.name, s.expected_date, s.magazine" + lim, params)
            else:
                cur.execute(
                    "SELECT l.code, l.name, l.category, r.magazine, s.periodicity, r.expected_date, r.status, "
                    "       r.marked_at, r.marked_by "
                    "FROM lib_receipts r JOIN lib_libraries l ON l.code=r.code "
                    "LEFT JOIN lib_schedule s ON s.code=r.code AND s.magazine=r.magazine AND s.expected_date=r.expected_date "
                    "WHERE TRUE" + extra + " ORDER BY l.name, r.expected_date, r.magazine" + lim, params)
            return status, cur.fetchall()
    finally:
        conn.close()


def _who(v):
    return "Admin" if v == "admin" else ("நூலகர்" if v else "")


@lib_bp.route("/api/admin/librarian/detail")
@admin_api
def api_admin_detail():
    status, rows = _detail_rows(limit=2000)
    ws, _we = week_bounds(today_ist())
    lo = ws - timedelta(weeks=CARRY_WEEKS)
    out = []
    for r in rows:
        mk = r["marked_at"]
        out.append({
            "magazine": r["magazine"], "periodicity": r["periodicity"] or "", "date": fmt_d(r["expected_date"]),
            "status": r["status"] or "",
            "markedAt": mk.astimezone(IST).strftime("%d/%m/%Y %H:%M") if mk else "",
            "markedDate": fmt_d(mk.astimezone(IST).date()) if mk else "",
            "markedBy": _who(r["marked_by"]),
            "delay": _delay_days(mk, r["expected_date"]) if r["status"] == "Received" else 0,
            "onLibrarianScreen": r["expected_date"] >= lo,
        })
    return jsonify({"success": True, "rows": out})


@lib_bp.route("/api/admin/librarian/export")
@admin_api
def api_admin_export():
    import xlsxwriter
    status, rows = _detail_rows()
    names = {"received": "Received", "notreceived": "Not received yet", "unregistered": "Not registered"}
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet(names[status])
    head = wb.add_format({"bold": True, "bg_color": "#1e4d8c", "font_color": "#ffffff", "border": 1})
    cols = ["Library Code", "Library", "Category", "Magazine", "Periodicity", "Expected Date"]
    if status != "unregistered":
        cols += ["Status", "Marked At", "Marked By"]
    if status == "received":
        cols += ["Delay (days)"]
    for c, h in enumerate(cols):
        ws.write(0, c, h, head)
    for i, r in enumerate(rows, start=1):
        vals = [r["code"], r["name"], r["category"], r["magazine"], r["periodicity"] or "", fmt_d(r["expected_date"])]
        if status != "unregistered":
            vals += [r["status"], r["marked_at"].astimezone(IST).strftime("%d/%m/%Y %H:%M") if r["marked_at"] else "",
                     _who(r["marked_by"])]
        if status == "received":
            vals += [_delay_days(r["marked_at"], r["expected_date"])]
        for c, v in enumerate(vals):
            ws.write(i, c, v)
    for c, w in enumerate([14, 38, 22, 34, 14, 14, 14, 18, 12, 12]):
        ws.set_column(c, c, w)
    ws.freeze_panes(1, 0)
    wb.close()
    fname = f"librarian_{status}_{today_ist().isoformat()}.xlsx"
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


# ---- Admin: ஒரு நூலகத்துக்காக Admin பதிவு செய்தல் ------------------------------------------------ #
def _weeks_arg(v, default=CARRY_WEEKS):
    """weeks=0 -> வரம்பின்றி (None)."""
    try:
        n = int(v)
    except (TypeError, ValueError):
        return default
    return None if n <= 0 else min(n, 520)


@lib_bp.route("/api/admin/librarian/libraries")
@admin_api
def api_admin_libraries():
    return _libraries_json()


@lib_bp.route("/api/admin/librarian/mark-list")
@admin_api
def api_admin_mark_list():
    code = (request.args.get("code") or "").strip()
    data = _library_week_payload(code, _weeks_arg(request.args.get("weeks")))
    if data is None:
        return jsonify({"success": False, "message": "நூலகம் கண்டறியப்படவில்லை."}), 404
    return jsonify(data)


@lib_bp.route("/api/admin/librarian/mark-save", methods=["POST"])
@admin_api
def api_admin_mark_save():
    body = request.get_json(force=True, silent=True) or {}
    code = (body.get("code") or "").strip()
    conn = _get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM lib_libraries WHERE code=%s", (code,))
            ok = cur.fetchone()
    finally:
        conn.close()
    if not ok:
        return jsonify({"success": False, "message": "நூலகம் கண்டறியப்படவில்லை."}), 404
    return _save_entries(code, body.get("entries") or [], None, "admin")


# =========================================================================== #
# HTML (தனித்து இயங்கும் பக்கங்கள்)
# =========================================================================== #
BASE_CSS = """
*{box-sizing:border-box} body{margin:0;font-family:'Noto Sans Tamil','Segoe UI',Latha,Tahoma,sans-serif;background:#eef2f9;color:#1b2540}
.top{background:linear-gradient(135deg,#0f2347,#1e4d8c);color:#fff;padding:14px 18px;display:flex;align-items:center;gap:12px;flex-wrap:wrap}
.top h1{font-size:18px;margin:0;flex:1;min-width:200px} .top a{color:#fff;text-decoration:none;background:rgba(255,255,255,.16);padding:7px 12px;border-radius:8px;font-size:13px}
.wrap{max-width:980px;margin:18px auto;padding:0 14px 90px}
.card{background:#fff;border-radius:14px;padding:18px;box-shadow:0 4px 18px rgba(15,35,71,.08);margin-bottom:16px}
label{display:block;font-size:13px;font-weight:600;color:#1e4d8c;margin:12px 0 6px}
select,input[type=date],input[type=file]{width:100%;padding:11px 12px;border:1.5px solid #c5cee3;border-radius:10px;font-size:15px;background:#fff;color:#111}
button,.btn{cursor:pointer;border:0;border-radius:10px;padding:11px 18px;font-size:15px;font-weight:700;background:#1e4d8c;color:#fff}
button:hover{background:#0f2347} button:disabled{opacity:.5;cursor:not-allowed}
.msg{margin-top:12px;padding:10px 12px;border-radius:10px;font-size:14px;display:none} .msg.ok{background:#e6f6ea;color:#1b6b34;display:block} .msg.err{background:#fdecea;color:#a12b20;display:block}
.muted{color:#5a6685;font-size:13px}
"""

SELECT_HTML = """<!DOCTYPE html><html lang="ta"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>நூலகத் தேர்வு</title>
<style>""" + BASE_CSS + """.wrap{max-width:520px}</style></head><body>
<div class="top"><h1>📚 நூலகர் நுழைவு</h1><a href="/logout">வெளியேறு</a></div>
<div class="wrap"><div class="card">
  <h2 style="margin:0 0 4px;font-size:18px">உங்கள் நூலகத்தைத் தேர்ந்தெடுக்கவும்</h2>
  <p class="muted" style="margin:0">முதலில் நூலக வகை, பிறகு நூலகத்தின் பெயர்.</p>
  <label for="cat">நூலக வகை</label>
  <select id="cat"><option value="">— தேர்ந்தெடுக்கவும் —</option></select>
  <label for="lib">நூலகத்தின் பெயர்</label>
  <select id="lib" disabled><option value="">— முதலில் வகையைத் தேர்ந்தெடுக்கவும் —</option></select>
  <button id="go" style="width:100%;margin-top:20px" disabled>தொடர்க ➜</button>
  <div id="msg" class="msg"></div>
</div></div>
<script>
let DATA=[];const cat=document.getElementById('cat'),lib=document.getElementById('lib'),go=document.getElementById('go'),msg=document.getElementById('msg');
function opt(v,t){const o=document.createElement('option');o.value=v;o.textContent=t;return o}
fetch('/api/librarian/libraries').then(r=>r.json()).then(d=>{
  if(!d.success){msg.className='msg err';msg.textContent=d.message||'பிழை';return}
  DATA=d.categories;DATA.forEach(c=>cat.appendChild(opt(c.category,c.category+' ('+c.libraries.length+')')));
}).catch(()=>{msg.className='msg err';msg.textContent='சேவையகத்தை அணுக முடியவில்லை.'});
cat.onchange=()=>{lib.innerHTML='';const c=DATA.find(x=>x.category===cat.value);
  if(!c){lib.appendChild(opt('','— முதலில் வகையைத் தேர்ந்தெடுக்கவும் —'));lib.disabled=true;go.disabled=true;return}
  lib.appendChild(opt('','— நூலகத்தைத் தேர்ந்தெடுக்கவும் —'));c.libraries.forEach(l=>lib.appendChild(opt(l.code,l.name)));lib.disabled=false;go.disabled=true};
lib.onchange=()=>{go.disabled=!lib.value};
go.onclick=async()=>{go.disabled=true;
  const r=await fetch('/api/librarian/select',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({code:lib.value})});
  const d=await r.json();if(d.success)location.href='/librarian/dashboard';else{msg.className='msg err';msg.textContent=d.message||'பிழை';go.disabled=false}};
</script></body></html>"""

DASH_HTML = """<!DOCTYPE html><html lang="ta"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>இதழ்கள் வருகை பதிவு</title>
<style>""" + BASE_CSS + """
.day{background:#1e4d8c;color:#fff;border-radius:10px;padding:9px 14px;margin:16px 0 8px;font-weight:700;font-size:16px}
.sec{border-radius:12px;padding:13px 16px;margin:24px 0 8px;font-weight:800;font-size:19px}
.sec.unreg{background:#fff3cd;color:#7a5600;border:1.5px solid #f0c75e}
.sec.pend{background:#fdecea;color:#a12b20;border:1.5px solid #f2b8b1}
.sec.this{background:#e8eefb;color:#0f2347;border:1.5px solid #b9c7e6}
.row{background:#fff;border-radius:12px;padding:12px 14px;margin-bottom:8px;box-shadow:0 2px 8px rgba(15,35,71,.07);display:flex;gap:12px;align-items:center;flex-wrap:wrap}
.row .nm{flex:1;min-width:180px;font-weight:700;font-size:16px} .row .nm small{display:block;color:#6b7694;font-weight:400;font-size:12px;margin-top:2px}
.opts{display:flex;gap:8px} .opts label{margin:0;display:block}
.opts input{display:none}
.opts span{display:inline-block;padding:9px 14px;border-radius:9px;border:1.5px solid #c5cee3;font-size:14px;font-weight:600;color:#44506f;background:#fff;cursor:pointer;user-select:none}
.opts input:checked + span.rc{background:#1b8a43;border-color:#1b8a43;color:#fff}
.opts input:checked + span.nr{background:#c0392b;border-color:#c0392b;color:#fff}
.bar{position:fixed;left:0;right:0;bottom:0;background:#fff;border-top:1px solid #d6ddee;padding:10px 14px;display:flex;gap:10px;align-items:center;justify-content:center;flex-wrap:wrap;z-index:5}
.bar .info{font-size:14px;color:#1b2540;font-weight:600} .bar .alt{background:#e8eefb;color:#1e4d8c}
.chip{display:inline-block;background:#e8eefb;color:#1e4d8c;border-radius:999px;padding:3px 10px;font-size:12px;margin-right:6px}
.top{padding:22px 24px}
.ttl{flex:1;min-width:260px;text-align:center;line-height:1.5}
.ttl .l1{font-size:clamp(22px,4.2vw,36px);font-weight:800} .ttl .l2{font-size:clamp(19px,3.4vw,29px);font-weight:700} .ttl .l3{font-size:clamp(19px,3.4vw,29px);font-weight:800;color:#ffe08a} .ttl .l4{font-size:clamp(18px,3.1vw,26px);font-weight:700;margin-top:4px;border-top:2px solid rgba(255,255,255,.4);padding-top:6px;display:inline-block}
.nav{display:flex;gap:8px;flex-wrap:wrap}
.tools{display:flex;gap:10px;align-items:center;flex-wrap:wrap}
.tools input[type=search]{flex:1;min-width:200px;padding:11px 12px;border:1.5px solid #c5cee3;border-radius:10px;font-size:15px}
.dl{display:none;gap:8px;margin-top:10px;flex-wrap:wrap} .dl.show{display:flex}
.dl .btn2{background:#e8eefb;color:#1e4d8c;border:1.5px solid #b9c7e6}
.boxes{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-bottom:16px}
@media(max-width:640px){.boxes{grid-template-columns:1fr}}
.bx{border-radius:14px;overflow:hidden;background:#fff;box-shadow:0 4px 18px rgba(15,35,71,.08)}
.bx h3{margin:0;padding:12px 14px;font-size:19px;color:#fff;line-height:1.35}
.bx.rc h3{background:#1b8a43} .bx.nr h3{background:#c0392b}
.mq{height:150px;overflow-y:auto;padding:6px 14px;scrollbar-width:thin}
.mq div.it{padding:7px 0;border-bottom:1px dashed #d6ddee;font-size:16px;font-weight:600}
.mq .none{color:#6b7694;font-size:13px;padding:10px 0}
.seg{display:flex;gap:8px;margin:0 0 12px;flex-wrap:wrap}
.seg button{background:#e8eefb;color:#1e4d8c;border:2px solid #b9c7e6;font-size:17px;padding:10px 20px}
.seg button.on{background:#1e4d8c;color:#fff;border-color:#1e4d8c}
.hint2{margin-top:8px;font-size:13px;color:#5a6685;width:100%}
#toast{position:fixed;left:50%;top:40%;transform:translate(-50%,-50%);background:#1b6b34;color:#fff;padding:16px 28px;border-radius:14px;font-size:18px;font-weight:800;box-shadow:0 8px 30px rgba(0,0,0,.3);z-index:50;display:none}
</style></head><body>
<div class="top"><div class="ttl"><div class="l1">தமிழ்நாடு அரசு பொது நூலகத்துறை</div><div class="l2">மாவட்ட நூலக ஆணைக்குழு, திண்டுக்கல்</div><div class="l3" id="libName">…</div><div class="l4">பருவ இதழ் வருகைப் பதிவேடு</div></div>
<div class="nav"><a href="/librarian/change">நூலகத்தை மாற்று</a><a href="/logout">வெளியேறு</a></div></div>
<div class="wrap">
  <div class="card" id="head"><div class="muted">ஏற்றுகிறது…</div></div>
  <div class="card">
    <button id="ordBtn" type="button">📋 உங்கள் நூலகத்திற்கு ஆணையிடப்பட்ட இதழ்கள்</button>
    <div class="dl" id="dl"><a class="btn btn2" href="/api/librarian/ordered-export?fmt=pdf">⬇️ PDF பதிவிறக்கம்</a><a class="btn btn2" href="/api/librarian/ordered-export?fmt=xlsx">⬇️ Excel பதிவிறக்கம்</a></div>
  </div>
  <div class="seg"><button type="button" id="mCur" class="on">📅 இந்த மாதம்</button><button type="button" id="mPrev">⏪ முந்தைய மாதம்</button></div>
  <div class="boxes">
    <div class="bx rc"><h3 id="rcT">🟢 இந்த மாதம் வந்த இதழ்கள்</h3><div class="mq" id="mqRc"></div></div>
    <div class="bx nr"><h3 id="nrT">🔴 இந்த மாதம் வராத இதழ்கள்</h3><div class="mq" id="mqNr"></div></div>
  </div>
  <div class="card tools"><input type="search" id="q" placeholder="🔍 இதழின் பெயரைத் தேடுக…" autocomplete="off"><div class="hint2">💡 "வரவில்லை" எனக் குறித்த இதழ் பின்னர் வந்தால்: இங்கே பெயரைத் தேடி, ✅ Received தேர்ந்து, 💾 சேமி அழுத்தவும்.</div></div>
  <div id="list"></div>
  <div id="msg" class="msg"></div>
</div>
<div id="toast"></div>
<div class="bar" id="bar" style="display:none">
  <span class="info" id="cnt"></span>
  <button class="alt" id="allRc" type="button">இந்த வாரம் — எல்லாவற்றையும் Received எனக் குறி</button>
  <button id="save" type="button">💾 சேமி</button>
</div>
<script>
const list=document.getElementById('list'),msg=document.getElementById('msg'),bar=document.getElementById('bar'),cnt=document.getElementById('cnt');
let ROWS=[];
function el(t,a,x){const e=document.createElement(t);if(a)for(const k in a)e.setAttribute(k,a[k]);if(x!==undefined)e.textContent=x;return e}
function update(){const n=ROWS.length,m=document.querySelectorAll('.opts input:checked').length;cnt.textContent='தேர்ந்தவை: '+m+' / '+n}
function fillBox(id,t,title,names){
  const b=document.getElementById(id);b.innerHTML='';b.scrollTop=0;document.getElementById(t).textContent=title+' ('+names.length+')';
  if(!names.length){b.appendChild(el('div',{class:'none'},'— இல்லை —'));return}
  const add=()=>names.forEach(n=>b.appendChild(el('div',{class:'it'},n)));add();
  if(b.scrollHeight>b.clientHeight+4){add();b.dataset.loop=b.scrollHeight/2}else b.dataset.loop=''}
const paused={mqRc:false,mqNr:false};
['mqRc','mqNr'].forEach(id=>{const b=document.getElementById(id);
  ['mouseenter','touchstart','focusin'].forEach(e=>b.addEventListener(e,()=>{paused[id]=true}));
  ['mouseleave','touchend','touchcancel','focusout'].forEach(e=>b.addEventListener(e,()=>{setTimeout(()=>{paused[id]=false},1500)}))});
setInterval(()=>{['mqRc','mqNr'].forEach(id=>{const b=document.getElementById(id),h=parseFloat(b.dataset.loop);
  if(!h||paused[id])return;b.scrollTop+=1;if(b.scrollTop>=h)b.scrollTop-=h})},45);
let MOFF=0;
async function loadBoxes(){
  try{const r=await fetch('/api/librarian/month?offset='+MOFF);const d=await r.json();if(!d.success)return;
    const p=MOFF?('முந்தைய மாதம் ('+d.label+')'):('இந்த மாதம் ('+d.label+')');
    const f=x=>x.magazine+' — '+x.dateLabel;
    fillBox('mqRc','rcT','🟢 '+p+' வந்த இதழ்கள்',d.received.map(f));
    fillBox('mqNr','nrT','🔴 '+p+' வராத இதழ்கள்',d.notReceived.map(f))}catch(e){}}
function setMonth(o){MOFF=o;document.getElementById('mCur').className=o?'':'on';document.getElementById('mPrev').className=o?'on':'';loadBoxes()}
function applyFilter(){
  const q=document.getElementById('q').value.trim().toLowerCase();
  document.querySelectorAll('#list .row').forEach(r=>{r.style.display=(!q||(r.getAttribute('data-n')||'').includes(q))?'':'none'});
  document.querySelectorAll('#list .day').forEach(dy=>{let n=dy.nextElementSibling,any=false;
    while(n&&n.classList.contains('row')){if(n.style.display!=='none')any=true;n=n.nextElementSibling}
    dy.style.display=(!q||any)?'':'none'});
  document.querySelectorAll('#list .sec').forEach(s=>{let n=s.nextElementSibling,any=false,hint=null;
    while(n&&!n.classList.contains('sec')&&n.tagName!=='DETAILS'){
      if(n.classList.contains('row')&&n.style.display!=='none')any=true;
      if(n.classList.contains('muted'))hint=n;n=n.nextElementSibling}
    s.style.display=(!q||any)?'':'none';if(hint)hint.style.display=(!q||any)?'':'none'});
  document.querySelectorAll('#list details').forEach(dt=>{
    const any=Array.from(dt.querySelectorAll('.row')).some(r=>r.style.display!=='none');
    dt.style.display=(!q||any)?'':'none';if(q&&any)dt.open=true})}
function toast(t){const e=document.getElementById('toast');e.textContent=t;e.style.display='block';clearTimeout(toast.t);toast.t=setTimeout(()=>{e.style.display='none'},2200)}
function render(d){
  const prevOpen=!!document.querySelector('#list details[open]');
  loadBoxes();
  const h=document.getElementById('head');h.innerHTML='';
  document.getElementById('libName').textContent=d.library.name||'';
  const s=el('div',{style:'margin-top:6px'});s.appendChild(el('span',{class:'chip'},d.library.category||''));s.appendChild(el('span',{class:'chip'},d.library.code||''));
  s.appendChild(el('span',{class:'chip'},'📅 இந்த வாரம்: '+d.weekStart+' – '+d.weekEnd));h.appendChild(s);
  const early=d.earlierUnregistered.length+d.earlierNotReceived.length;
  if(early)h.appendChild(el('div',{class:'muted',style:'margin-top:8px'},'முந்தைய வாரங்களின் இதழ்கள் அதிகபட்சம் '+d.carryWeeks+' வாரங்கள் வரை இங்கே தெரியும்.'));
  ROWS=[];list.innerHTML='';
  function addRows(parent,rows,k,early){
    let last='';
    rows.forEach(r=>{
      if(r.date!==last){last=r.date;parent.appendChild(el('div',{class:'day'},r.day+' — '+r.dateLabel+(r.date===d.today?'  (இன்று)':'')+(early&&r.weeksAgo?'  · '+r.weeksAgo+' வாரத்துக்கு முன்':'')))}
      const i=ROWS.push({magazine:r.magazine,date:r.date})-1;
      const row=el('div',{class:'row','data-n':r.magazine.toLowerCase()});const nm=el('div',{class:'nm'},r.magazine);if(r.periodicity)nm.appendChild(el('small',null,r.periodicity));row.appendChild(nm);
      const o=el('div',{class:'opts'});
      [['Received','rc','✅ Received'],['Not received','nr','❌ Not received']].forEach(([v,c,t])=>{
        const l=el('label');const inp=el('input',{type:'radio',name:'r'+i,value:v,'data-sec':k});if(r.status===v)inp.checked=true;inp.onchange=update;
        l.appendChild(inp);l.appendChild(el('span',{class:c},t));o.appendChild(l)});
      row.appendChild(o);parent.appendChild(row)})
  }
  const openRows=d.rows.filter(r=>r.status!=='Received'),doneRows=d.rows.filter(r=>r.status==='Received');
  [['⚠️ கடந்த வாரங்களில் பதிவு செய்யாதவை','unreg',d.earlierUnregistered,'இவை இன்னும் பதிவு செய்யப்படவில்லை — Received / Not received தேர்ந்து சேமிக்கவும்.',true],
   ['⏳ முந்தைய வாரங்களில் இன்னும் வராதவை','pend',d.earlierNotReceived,'முன்பு "Not received" எனக் குறித்தவை — இதழ் வந்திருந்தால் Received எனக் குறிக்கவும்; வராவிட்டால் அப்படியே விடவும்.',true],
   ['📅 இந்த வாரம் — இன்னும் பதிவாகாதவை / வராதவை ('+d.weekStart+' – '+d.weekEnd+')','this',openRows,'',false]].forEach(([title,k,rows,hint,early])=>{
    if(!rows.length)return;
    list.appendChild(el('div',{class:'sec '+k},title+' ('+rows.length+')'));
    if(hint)list.appendChild(el('div',{class:'muted',style:'margin:0 2px 4px'},hint));
    addRows(list,rows,k,early)});
  if(!d.rows.length)list.appendChild(el('div',{class:'card muted',style:'margin-top:14px'},'இந்த வாரத்துக்கு ('+d.weekStart+' – '+d.weekEnd+') வர வேண்டிய இதழ்கள் எதுவும் இல்லை.'));
  else if(!openRows.length)list.appendChild(el('div',{class:'card',style:'margin-top:14px;background:#e6f6ea;color:#1b6b34;font-weight:700'},'✅ இந்த வாரத்துக்கான எல்லா இதழ்களும் Received எனப் பதிவாகிவிட்டன.'));
  if(doneRows.length){
    const det=el('details',{style:'margin-top:18px'});
    det.appendChild(el('summary',{style:'cursor:pointer;font-weight:700;color:#1b6b34;padding:8px 2px'},'✅ இந்த வாரம் Received எனப் பதிவானவை ('+doneRows.length+') — மாற்ற வேண்டுமானால் திறக்கவும்'));
    addRows(det,doneRows,'this',false);if(prevOpen)det.open=true;list.appendChild(det)}
  if(!ROWS.length){bar.style.display='none';return}
  bar.style.display='flex';update();applyFilter();
}
async function load(){
  const r=await fetch('/api/librarian/week');const d=await r.json();
  if(d.needLibrary){location.href='/librarian/select';return}
  if(!d.success){msg.className='msg err';msg.textContent=d.message||'பிழை';return}
  render(d)}
document.getElementById('q').addEventListener('input',applyFilter);
document.getElementById('mCur').onclick=()=>setMonth(0);
document.getElementById('mPrev').onclick=()=>setMonth(1);
document.getElementById('ordBtn').onclick=()=>document.getElementById('dl').classList.toggle('show');
document.getElementById('allRc').onclick=()=>{document.querySelectorAll('.opts input[value="Received"][data-sec="this"]').forEach(i=>i.checked=true);update()};
document.getElementById('save').onclick=async()=>{
  const entries=[];ROWS.forEach((r,i)=>{const c=document.querySelector('input[name="r'+i+'"]:checked');if(c)entries.push({magazine:r.magazine,date:r.date,status:c.value})});
  if(!entries.length){msg.className='msg err';msg.textContent='குறைந்தது ஒரு இதழுக்காவது Received / Not received தேர்ந்தெடுக்கவும்.';return}
  const b=document.getElementById('save');b.disabled=true;
  try{const r=await fetch('/api/librarian/save',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({entries})});
    const d=await r.json();const y=window.scrollY;if(d.success){msg.className='msg';msg.textContent='';await load();window.scrollTo(0,y);toast('✅ சேமிக்கப்பட்டது')}else{msg.className='msg err';msg.textContent=d.message||'பிழை'}}
  catch(e){msg.className='msg err';msg.textContent='சேமிக்க முடியவில்லை: '+e}finally{b.disabled=false}};
load();
</script></body></html>"""

ADMIN_HTML = """<!DOCTYPE html><html lang="ta"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>நூலகர் பதிவு — Admin</title>
<style>""" + BASE_CSS + """
.wrap{max-width:1100px} h2{margin:0 0 8px;font-size:17px;color:#0f2347}
.tabs{display:flex;gap:8px;margin-bottom:12px;flex-wrap:wrap} .tab{background:#e8eefb;color:#1e4d8c;flex:1;min-width:150px} .tab.on{background:#1e4d8c;color:#fff}
.flt{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:10px;align-items:end}
table{width:100%;border-collapse:collapse;font-size:14px} th{background:#1e4d8c;color:#fff;text-align:left;padding:8px 10px;position:sticky;top:0}
td{padding:8px 10px;border-bottom:1px solid #e3e8f4;vertical-align:top} tr.lib{cursor:pointer} tr.lib:hover{background:#f3f6fd}
.sub{background:#f8faff} .sub table th{background:#5b7fb8}
.ok-c{color:#1b6b34;font-weight:700} .bad-c{color:#c0392b;font-weight:700} .warn-c{color:#b26a00;font-weight:700}
.kpi{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0} .kpi div{background:#e8eefb;border-radius:10px;padding:8px 14px;font-size:14px}
.sm{padding:6px 12px;font-size:13px} .danger{background:#c0392b} .alt{background:#e8eefb;color:#1e4d8c}
.scroll{overflow-x:auto} details summary{cursor:pointer;font-weight:700;color:#1e4d8c}
.mh{border-radius:10px;padding:9px 12px;margin:16px 0 6px;font-weight:800;font-size:14px}
.mh.unreg{background:#fff3cd;color:#7a5600} .mh.pend{background:#fdecea;color:#a12b20} .mh.this{background:#e8eefb;color:#0f2347}
.md{background:#1e4d8c;color:#fff;border-radius:8px;padding:6px 10px;margin:10px 0 6px;font-weight:700;font-size:13px}
.mr{background:#fff;border:1px solid #e3e8f4;border-radius:10px;padding:10px 12px;margin-bottom:6px;display:flex;gap:12px;align-items:center;flex-wrap:wrap}
.mr .nm{flex:1;min-width:180px;font-weight:600} .mr .nm small{display:block;color:#6b7694;font-weight:400;font-size:12px}
.opts{display:flex;gap:8px} .opts label{margin:0;display:block} .opts input{display:none}
.opts span{display:inline-block;padding:8px 13px;border-radius:9px;border:1.5px solid #c5cee3;font-size:14px;font-weight:600;color:#44506f;background:#fff;cursor:pointer;user-select:none}
.opts input:checked + span.rc{background:#1b8a43;border-color:#1b8a43;color:#fff}
.opts input:checked + span.nr{background:#c0392b;border-color:#c0392b;color:#fff}
.mbar{position:sticky;bottom:0;background:#fff;border-top:1px solid #d6ddee;padding:10px 0;margin-top:12px;display:none;gap:10px;align-items:center;flex-wrap:wrap}
.chip{display:inline-block;background:#e8eefb;color:#1e4d8c;border-radius:999px;padding:3px 10px;font-size:12px;margin-right:6px}
</style></head><body>
<div class="top"><h1>🧑‍💼 நூலகர் பதிவு (இதழ் வருகை)</h1><a href="/">← முகப்பு</a></div>
<div class="wrap">

<div class="card">
  <h2>① Quarter Excel ஏற்றுதல் (Q3 / Q4 …)</h2>
  <p class="muted" style="margin:0 0 6px">நெடுவரிசை வரிசை: <b>A – Library code, B – Magazine Name, C – Periodicity, D – Expected Date</b> (தலைப்பு வரிசை இருக்கலாம்). கோப்பு .xlsx / .xls / .csv. Quarter தேதியிலிருந்து தானாகக் கணிக்கப்படும்.</p>
  <label>Excel கோப்பு</label><input type="file" id="qfile" accept=".xlsx,.xls,.csv">
  <label style="display:flex;gap:8px;align-items:center;font-weight:500;color:#1b2540"><input type="checkbox" id="qreplace" checked style="width:auto"> இந்தக் கோப்பிலுள்ள Quarter-ன் பழைய அட்டவணையை மாற்று (நூலகர் பதிவுகள் அழியாது)</label>
  <button id="qgo" style="margin-top:12px">⬆️ ஏற்று</button>
  <div id="qmsg" class="msg"></div>
  <div class="scroll" style="margin-top:14px"><table id="qtab"><thead><tr><th>Quarter</th><th>வரிசைகள்</th><th>தேதிகள்</th><th>நூலகங்கள்</th><th>இதழ்கள்</th><th>பதிவானவை</th><th></th></tr></thead><tbody></tbody></table></div>
</div>

<div class="card">
  <details><summary>நூலகப் பட்டியல் (name, category, code) புதுப்பிக்க</summary>
    <p class="muted">தற்போது <b id="libcnt">?</b> நூலகங்கள் உள்ளன. புதிய நூலகம் சேர்க்க/பெயர் மாற்ற, name · category · code உள்ள கோப்பை ஏற்றவும் (code வாரியாகப் புதுப்பிக்கப்படும்).</p>
    <input type="file" id="lfile" accept=".xlsx,.xls,.csv"><button id="lgo" class="sm" style="margin-top:10px">⬆️ ஏற்று</button><div id="lmsg" class="msg"></div>
  </details>
</div>

<div class="card">
  <h2>② நூலகர் பதிவு நிலை</h2>
  <div class="tabs"><button class="tab on" id="t1">✅ பதிவு செய்தவை (Received)</button><button class="tab" id="t2">⏳ இன்னும் வராதவை</button><button class="tab" id="t3">⚠️ பதிவே செய்யாதவை</button></div>
  <div class="flt">
    <div><label>நூலக வகை</label><select id="fcat"><option value="">எல்லாம்</option></select></div>
    <div><label>தேதி முதல்</label><input type="date" id="ffrom"></div>
    <div><label>தேதி வரை</label><input type="date" id="fto"></div>
    <div style="display:flex;gap:8px"><button id="fgo">🔍 காட்டு</button><button id="fweek" class="alt">இந்த வாரம்</button></div>
  </div>
  <div class="kpi" id="kpi"></div>
  <div style="margin:6px 0 10px"><button id="fxl" class="sm alt">📥 Excel பதிவிறக்கு</button> <span class="muted">வரிசையைக் கிளிக் செய்தால் இதழ்கள் விரியும் · "✏️ பதிவு" பொத்தான் அந்த நூலகத்துக்கு Admin பதிவுக்குக் கொண்டு செல்லும்</span></div>
  <div class="scroll"><table id="stab"><thead></thead><tbody></tbody></table></div>
</div>

<div class="card" id="mcard">
  <h2>③ நூலகத்துக்காக Admin பதிவு (Received / Not received)</h2>
  <p class="muted" style="margin:0 0 6px">ஒரு நூலகத்தைத் தேர்ந்தெடுத்து, நூலகருக்குப் பதிலாக நீங்களே பதிவு செய்யலாம். இந்த வாரம் மற்றும் அதற்கு முந்தைய தேதிகள் மட்டும்; எதிர்காலத் தேதிகள் முடியாது. இங்கே பதிவானவை "Admin" எனக் குறிக்கப்படும்.</p>
  <div class="flt">
    <div><label>நூலக வகை</label><select id="mcat"><option value="">— தேர்ந்தெடுக்கவும் —</option></select></div>
    <div><label>நூலகத்தின் பெயர்</label><select id="mlib"><option value="">— நூலகம் —</option></select></div>
    <div><label>முந்தைய வாரங்கள்</label><select id="mweeks"><option value="4">4 வாரம்</option><option value="8">8 வாரம்</option><option value="13">13 வாரம்</option><option value="26">26 வாரம்</option><option value="0">எல்லாம்</option></select></div>
    <div><button id="mload">📋 இதழ்களைக் காட்டு</button></div>
  </div>
  <div id="mmsg" class="msg"></div>
  <div id="mhead" style="margin-top:12px"></div>
  <div id="mlist"></div>
  <div class="mbar" id="mbar"><span id="mcnt" style="font-weight:600"></span><button id="mallrc" class="alt" type="button">இந்த வாரம் — எல்லாம் Received</button><button id="msave" type="button">💾 சேமி (Admin)</button></div>
</div>
</div>
<script>
let MODE='received',OV=null,LIBS=[],MROWS=[],MCODE='';
const $=id=>document.getElementById(id);
function el(t,a,x){const e=document.createElement(t);if(a)for(const k in a)e.setAttribute(k,a[k]);if(x!==undefined)e.textContent=x;return e}
function td(t,cls){const c=el('td',cls?{class:cls}:null,t===undefined?'':String(t));return c}
function say(id,ok,t,warn){const m=$(id);m.className='msg '+(ok?'ok':'err');m.textContent=t;(warn||[]).forEach(w=>{m.appendChild(el('div',{style:'margin-top:4px'},'⚠️ '+w))})}
async function api(u,o){const r=await fetch(u,o);if(r.status===401){location.href='/login';throw 0}return r.json()}

async function loadOverview(){
  OV=await api('/api/admin/librarian/overview');
  $('libcnt').textContent=OV.libraries;
  const fc=$('fcat');fc.length=1;OV.categories.forEach(c=>{const o=el('option',{value:c},c);fc.appendChild(o)});
  const tb=$('qtab').querySelector('tbody');tb.innerHTML='';
  if(!OV.quarters.length){const tr=el('tr');const c=td('இன்னும் எந்த Quarter-ம் ஏற்றப்படவில்லை.','muted');c.colSpan=7;tr.appendChild(c);tb.appendChild(tr)}
  OV.quarters.forEach(q=>{const tr=el('tr');[q.quarter,q.rows,q.from+' – '+q.to,q.libs,q.mags,q.recorded].forEach(v=>tr.appendChild(td(v)));
    const c=el('td');const b=el('button',{class:'sm danger',type:'button'},'🗑️ நீக்கு');
    b.onclick=async()=>{if(!confirm(q.quarter+' அட்டவணையையும் அதற்கான '+q.recorded+' நூலகர் பதிவுகளையும் நீக்கவா?'))return;
      const d=await api('/api/admin/librarian/delete-quarter',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({quarter:q.quarter})});
      say('qmsg',d.success,d.message);loadOverview();loadSummary()};
    c.appendChild(b);tr.appendChild(c);tb.appendChild(tr)})
}
$('qgo').onclick=async()=>{
  const f=$('qfile').files[0];if(!f){say('qmsg',false,'கோப்பைத் தேர்ந்தெடுக்கவும்.');return}
  const fd=new FormData();fd.append('file',f);fd.append('mode',$('qreplace').checked?'replace':'add');
  $('qgo').disabled=true;say('qmsg',true,'ஏற்றுகிறது… (பெரிய கோப்புக்கு சில விநாடிகள் ஆகலாம்)');
  try{const d=await api('/api/admin/librarian/upload-quarter',{method:'POST',body:fd});say('qmsg',d.success,d.message,d.warnings);if(d.success){$('qfile').value='';loadOverview();loadSummary()}}
  catch(e){say('qmsg',false,'ஏற்ற முடியவில்லை (நேரம் முடிந்திருக்கலாம்): '+e)}finally{$('qgo').disabled=false}};
$('lgo').onclick=async()=>{
  const f=$('lfile').files[0];if(!f){say('lmsg',false,'கோப்பைத் தேர்ந்தெடுக்கவும்.');return}
  const fd=new FormData();fd.append('file',f);const d=await api('/api/admin/librarian/upload-libraries',{method:'POST',body:fd});
  say('lmsg',d.success,d.message);if(d.success){loadOverview();loadLibs()}};

function qs(extra){const p=new URLSearchParams({status:MODE});if($('fcat').value)p.set('category',$('fcat').value);
  if($('ffrom').value)p.set('from',$('ffrom').value);if($('fto').value)p.set('to',$('fto').value);for(const k in (extra||{}))p.set(k,extra[k]);return p.toString()}
const HEADS={
  received:['நூலகம்','வகை','Received','தாமதமாகப் பதிவானவை','அதிக தாமதம் (நாள்)','Admin பதிவு','கடைசிப் பதிவு',''],
  notreceived:['நூலகம்','வகை','இன்னும் வராதவை','இதழ்கள்','பழைய தேதி','கடைசிப் பதிவு',''],
  unregistered:['நூலகம்','வகை','பதிவு செய்யாதவை','இதழ்கள்','முதல் தேதி','இதுவரை பதிவு செய்தவை','']};
async function loadSummary(){
  const d=await api('/api/admin/librarian/summary?'+qs());
  const th=$('stab').querySelector('thead'),tb=$('stab').querySelector('tbody');th.innerHTML='';tb.innerHTML='';
  const hr=el('tr');HEADS[MODE].forEach(h=>hr.appendChild(el('th',null,h)));th.appendChild(hr);
  const k=$('kpi');k.innerHTML='';
  if(MODE==='received'){k.appendChild(el('div',null,'Received பதிவு செய்த நூலகங்கள்: '+d.items.length+' / '+d.totalLibraries));k.appendChild(el('div',null,'மொத்த Received: '+d.totalRows));k.appendChild(el('div',null,'தாமதமாகப் பதிவானவை: '+d.lateTotal))}
  else if(MODE==='notreceived'){k.appendChild(el('div',null,'இன்னும் வராத இதழ்கள் உள்ள நூலகங்கள்: '+d.items.length+' / '+d.totalLibraries));k.appendChild(el('div',null,'வராத இதழ் வரிசைகள்: '+d.totalRows))}
  else{k.appendChild(el('div',null,'பதிவு செய்யாத பதிவுகள் உள்ள நூலகங்கள்: '+d.items.length+' / '+d.totalLibraries));k.appendChild(el('div',null,'பதிவு செய்யப்படாத வரிசைகள் (இன்று வரை): '+d.totalRows));
    const never=d.items.filter(i=>i.recordedTotal===0).length;k.appendChild(el('div',null,'இதுவரை ஒன்றுமே பதிவு செய்யாத நூலகங்கள்: '+never))}
  if(!d.items.length){const tr=el('tr');const c=td('பதிவுகள் இல்லை.','muted');c.colSpan=8;tr.appendChild(c);tb.appendChild(tr)}
  d.items.forEach(i=>{
    const tr=el('tr',{class:'lib'});tr.appendChild(td(i.name));tr.appendChild(td(i.category));
    if(MODE==='received'){tr.appendChild(td(i.count,'ok-c'));tr.appendChild(td(i.late,i.late?'warn-c':''));tr.appendChild(td(i.maxLate||''));tr.appendChild(td(i.byAdmin||''));tr.appendChild(td(i.lastAt))}
    else if(MODE==='notreceived'){tr.appendChild(td(i.count,'bad-c'));tr.appendChild(td(i.magazines));tr.appendChild(td(i.firstDue));tr.appendChild(td(i.lastAt))}
    else{tr.appendChild(td(i.count,'bad-c'));tr.appendChild(td(i.magazines));tr.appendChild(td(i.firstDue));tr.appendChild(td(i.recordedTotal))}
    const ac=el('td');const mb=el('button',{class:'sm alt',type:'button'},'✏️ பதிவு');mb.onclick=e=>{e.stopPropagation();openMark(i.code,i.category)};ac.appendChild(mb);tr.appendChild(ac);
    let open=false,sub=null;
    tr.onclick=async()=>{ if(open){sub.remove();open=false;return}
      open=true;sub=el('tr',{class:'sub'});const c=el('td');c.colSpan=8;c.textContent='ஏற்றுகிறது…';sub.appendChild(c);tr.after(sub);
      const r=await api('/api/admin/librarian/detail?'+qs({code:i.code}));c.textContent='';
      const t=el('table');const h=el('tr');
      (MODE==='received'?['இதழ்','Periodicity','எதிர்பார்த்தது → வந்தது','பதிவு நேரம்','பதிவு செய்தவர்']:MODE==='notreceived'?['தேதி','இதழ்','Periodicity','பதிவு நேரம்','பதிவு செய்தவர்','நூலகர் திரையில்']:['தேதி','இதழ்','Periodicity','நூலகர் திரையில்']).forEach(x=>h.appendChild(el('th',null,x)));t.appendChild(h);
      r.rows.forEach(w=>{const rr=el('tr');
        if(MODE==='received'){rr.appendChild(td(w.magazine));rr.appendChild(td(w.periodicity));
          rr.appendChild(td(w.date+' → '+(w.markedDate||'')+(w.delay>0?' ('+w.delay+' நாள் தாமதம்)':''),w.delay>0?'warn-c':'ok-c'));rr.appendChild(td(w.markedAt));rr.appendChild(td(w.markedBy))}
        else if(MODE==='notreceived'){rr.appendChild(td(w.date));rr.appendChild(td(w.magazine));rr.appendChild(td(w.periodicity));rr.appendChild(td(w.markedAt));rr.appendChild(td(w.markedBy));rr.appendChild(td(w.onLibrarianScreen?'✔ தெரியும்':'— (4 வாரத்துக்கு மேல்)'))}
        else{rr.appendChild(td(w.date));rr.appendChild(td(w.magazine));rr.appendChild(td(w.periodicity));rr.appendChild(td(w.onLibrarianScreen?'✔ தெரியும்':'— (4 வாரத்துக்கு மேல்)'))}
        t.appendChild(rr)});
      c.appendChild(t);if(r.rows.length>=2000)c.appendChild(el('div',{class:'muted'},'முதல் 2000 வரிசைகள் மட்டும் — முழுமைக்கு Excel பதிவிறக்கவும்.'))};
    tb.appendChild(tr)})
}
function setMode(m){MODE=m;$('t1').classList.toggle('on',m==='received');$('t2').classList.toggle('on',m==='notreceived');$('t3').classList.toggle('on',m==='unregistered');loadSummary()}
$('t1').onclick=()=>setMode('received');$('t2').onclick=()=>setMode('notreceived');$('t3').onclick=()=>setMode('unregistered');
$('fgo').onclick=loadSummary;
$('fweek').onclick=()=>{$('ffrom').value=OV.weekStart;$('fto').value=OV.weekEnd;loadSummary()};
$('fxl').onclick=()=>{location.href='/api/admin/librarian/export?'+qs()};

/* ---------- ③ Admin பதிவு ---------- */
async function loadLibs(){
  const d=await api('/api/admin/librarian/libraries');LIBS=d.categories||[];
  const c=$('mcat');c.length=1;LIBS.forEach(x=>c.appendChild(el('option',{value:x.category},x.category+' ('+x.libraries.length+')')));fillLibs()}
function fillLibs(){const l=$('mlib');l.length=1;const c=LIBS.find(x=>x.category===$('mcat').value);if(c)c.libraries.forEach(x=>l.appendChild(el('option',{value:x.code},x.name)))}
$('mcat').onchange=fillLibs;
function mcount(){$('mcnt').textContent='தேர்ந்தவை: '+document.querySelectorAll('#mlist .opts input:checked').length+' / '+MROWS.length}
function renderMark(d){
  const h=$('mhead');h.innerHTML='';h.appendChild(el('span',{style:'font-weight:700;font-size:16px;margin-right:8px'},d.library.name||''));
  h.appendChild(el('span',{class:'chip'},d.library.category||''));h.appendChild(el('span',{class:'chip'},d.library.code||''));
  const box=$('mlist');box.innerHTML='';MROWS=[];
  [['⚠️ கடந்த வாரங்களில் பதிவு செய்யாதவை','unreg',d.earlierUnregistered],['⏳ முந்தைய வாரங்களில் இன்னும் வராதவை','pend',d.earlierNotReceived],['📅 இந்த வாரம் ('+d.weekStart+' – '+d.weekEnd+')','this',d.rows]].forEach(([t,k,rows])=>{
    if(!rows.length)return;box.appendChild(el('div',{class:'mh '+k},t+' ('+rows.length+')'));let last='';
    rows.forEach(r=>{
      if(r.date!==last){last=r.date;box.appendChild(el('div',{class:'md'},r.day+' — '+r.dateLabel+(r.date===d.today?'  (இன்று)':'')))}
      const i=MROWS.push({magazine:r.magazine,date:r.date})-1;
      const row=el('div',{class:'mr'});const nm=el('div',{class:'nm'},r.magazine);if(r.periodicity)nm.appendChild(el('small',null,r.periodicity));row.appendChild(nm);
      const o=el('div',{class:'opts'});
      [['Received','rc','✅ Received'],['Not received','nr','❌ Not received']].forEach(([v,c,tt])=>{
        const l=el('label');const inp=el('input',{type:'radio',name:'m'+i,value:v,'data-sec':k});if(r.status===v)inp.checked=true;inp.onchange=mcount;
        l.appendChild(inp);l.appendChild(el('span',{class:c},tt));o.appendChild(l)});
      row.appendChild(o);box.appendChild(row)})});
  if(!MROWS.length)box.appendChild(el('div',{class:'muted',style:'margin-top:10px'},'இந்தக் காலத்துக்குப் பதிவு செய்ய இதழ்கள் எதுவும் இல்லை.'));
  $('mbar').style.display=MROWS.length?'flex':'none';mcount()}
async function loadMark(){
  const code=$('mlib').value;if(!code){say('mmsg',false,'முதலில் நூலகத்தைத் தேர்ந்தெடுக்கவும்.');return}
  MCODE=code;const d=await api('/api/admin/librarian/mark-list?code='+encodeURIComponent(code)+'&weeks='+$('mweeks').value);
  if(!d.success){say('mmsg',false,d.message||'பிழை');return}
  renderMark(d)}
$('mload').onclick=()=>{$('mmsg').className='msg';loadMark()};
function openMark(code,cat){$('mcat').value=cat;fillLibs();$('mlib').value=code;$('mmsg').className='msg';loadMark().then(()=>$('mcard').scrollIntoView({behavior:'smooth'}))}
$('mallrc').onclick=()=>{document.querySelectorAll('#mlist .opts input[value="Received"][data-sec="this"]').forEach(i=>i.checked=true);mcount()};
$('msave').onclick=async()=>{
  const entries=[];MROWS.forEach((r,i)=>{const c=document.querySelector('#mlist input[name="m'+i+'"]:checked');if(c)entries.push({magazine:r.magazine,date:r.date,status:c.value})});
  if(!entries.length){say('mmsg',false,'குறைந்தது ஒரு இதழுக்காவது Received / Not received தேர்ந்தெடுக்கவும்.');return}
  $('msave').disabled=true;
  try{const d=await api('/api/admin/librarian/mark-save',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({code:MCODE,entries})});
    say('mmsg',d.success,d.message||'');if(d.success){await loadMark();loadSummary()}}
  catch(e){say('mmsg',false,'சேமிக்க முடியவில்லை: '+e)}finally{$('msave').disabled=false}};
loadOverview().then(loadSummary);loadLibs();
</script></body></html>"""
